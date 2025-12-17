import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from testDjangosite.settings import BASE_DIR


def create_plot_description_excel(data, wb=None, save: bool = True):
    if save:
        output_filename = f'{BASE_DIR}/media/excel_files/desc/desc_{data["id"]}.xlsx'

    if wb is None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Описание участка"
    else:
        ws = wb.create_sheet("Описание участка")

    bold_font = Font(bold=True)
    border = Border(left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin'))
    center_align = Alignment(horizontal='center', vertical='center')

    wrap_align = Alignment(horizontal='left', vertical='top', wrap_text=True)

    column_widths = {
        'A': 30, 'B': 40, 'C': 25, 'D': 20, 'E': 20, 'F': 20, 'G': 20
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    current_row = 1

    def safe_get_value(key, default=''):
        value = data.get(key, default)
        if isinstance(value, dict):
            return str(value)
        return value

    def create_wrapped_cell(ws, cell_range, value, is_bold=False, is_merged=True):
        if is_merged:
            ws.merge_cells(cell_range)

        cell = ws[cell_range.split(':')[0]]
        cell.value = value
        cell.alignment = wrap_align

        if is_bold:
            cell.font = bold_font

        if value:
            text = str(value)
            lines = text.split('\n')
            line_count = len(lines)

            for line in lines:
                if len(line) > 50:
                    line_count += len(line) // 50

            height = max(20, line_count * 15)
            row_num = cell.row
            ws.row_dimensions[row_num].height = height

        return cell

    title_cell = create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                                     f"ОПИСАНИЕ УЧАСТКА № {safe_get_value('number_region')}",
                                     is_bold=True)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = center_align
    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        "Способ лесовосстановления: (искусственный, комбинированный, естественный вследствие мер содействия лесовосстановлению, естественный вследствие природных процессов):",
                        is_bold=True)
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        safe_get_value('method_of_reforestation'))
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        "Метод лесовосстановления (сохранение подроста, минерализация почвы, оставление семенных деревьев, др.):",
                        is_bold=True)
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        safe_get_value('recovery_method'))
    current_row += 1

    ws[f'A{current_row}'] = "Год отнесения к землям, на которых расположены леса,"
    ws[f'A{current_row}'].font = bold_font
    ws[f'A{current_row}'].alignment = wrap_align
    ws[f'B{current_row}'] = safe_get_value('year_assignment_land')
    ws[f'B{current_row}'].alignment = wrap_align
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        "Местонахождение участка:",
                        is_bold=True)
    current_row += 1

    ws[f'A{current_row}'] = "Субъект РФ:"
    ws[f'A{current_row}'].font = bold_font
    ws[f'A{current_row}'].alignment = wrap_align
    ws[f'B{current_row}'] = safe_get_value('id_subject_rf')
    ws[f'B{current_row}'].alignment = wrap_align
    current_row += 1

    ws[f'A{current_row}'] = "Лесничество:"
    ws[f'A{current_row}'].font = bold_font
    ws[f'A{current_row}'].alignment = wrap_align
    ws[f'B{current_row}'] = safe_get_value('forestly')
    ws[f'B{current_row}'].alignment = wrap_align
    current_row += 1

    ws[f'A{current_row}'] = "Участковое лесничество:"
    ws[f'A{current_row}'].font = bold_font
    ws[f'A{current_row}'].alignment = wrap_align
    ws[f'B{current_row}'] = safe_get_value('district_forestly')
    ws[f'B{current_row}'].alignment = wrap_align
    current_row += 1

    ws[f'A{current_row}'] = "Урочище"
    ws[f'A{current_row}'].font = bold_font
    ws[f'A{current_row}'].alignment = wrap_align
    ws[f'B{current_row}'] = safe_get_value('name_dacha')
    ws[f'B{current_row}'].alignment = wrap_align
    current_row += 1

    ws[f'A{current_row}'] = "Квартал"
    ws[f'A{current_row}'].font = bold_font
    ws[f'A{current_row}'].alignment = wrap_align
    ws[f'B{current_row}'] = safe_get_value('name_quarter')
    ws[f'B{current_row}'].alignment = wrap_align

    ws[f'C{current_row}'] = "Выдел"
    ws[f'C{current_row}'].font = bold_font
    ws[f'C{current_row}'].alignment = wrap_align
    ws[f'D{current_row}'] = safe_get_value('soil_lot')
    ws[f'D{current_row}'].alignment = wrap_align

    ws[f'E{current_row}'] = "Площадь, га"
    ws[f'E{current_row}'].font = bold_font
    ws[f'E{current_row}'].alignment = wrap_align
    ws[f'F{current_row}'] = safe_get_value('sample_region')
    ws[f'F{current_row}'].alignment = wrap_align
    current_row += 1

    ws[f'A{current_row}'] = "Год образования категории фонда лесовосстановления (вырубки, гари)"
    ws[f'A{current_row}'].font = bold_font
    ws[f'A{current_row}'].alignment = wrap_align
    ws[f'B{current_row}'] = safe_get_value('year_format_fond_trees')
    ws[f'B{current_row}'].alignment = wrap_align
    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        "данные о проведенных мероприятиях по уходу за лесами на участке (вид и год проведения)",
                        is_bold=True)
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        safe_get_value('inf_restore_forest'))
    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        "Породный состав молодняка по Акту отнесения земель, предназначенных для лесовосстановления, к землям, на которых, расположены леса,",
                        is_bold=True)
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        safe_get_value('breed_structure_sapling_act_land'))
    current_row += 1

    ws[f'A{current_row}'] = "хозяйство по Акту отнесения земель"
    ws[f'A{current_row}'].font = bold_font
    ws[f'A{current_row}'].alignment = wrap_align
    ws[f'B{current_row}'] = safe_get_value('economy_act_land')
    ws[f'B{current_row}'].alignment = wrap_align
    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        "Породный состав молодняка по данным натурного обследования",
                        is_bold=True)
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        safe_get_value('breed_composition_sapling_data_surver'))
    current_row += 1

    ws[f'A{current_row}'] = "хозяйство по данным натурного обследования"
    ws[f'A{current_row}'].font = bold_font
    ws[f'A{current_row}'].alignment = wrap_align
    ws[f'B{current_row}'] = safe_get_value('farm_according_data_survey')
    ws[f'B{current_row}'].alignment = wrap_align
    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        "Изменение породного и качественного состава молодняка после отнесения участка к землям, на которых расположены леса на момент натурного обследования лесного участка",
                        is_bold=True)
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        safe_get_value('change_breed_and_structure_sapling'))
    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        "Вывод по результатам обследования",
                        is_bold=True)
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        safe_get_value('results_surtvey'))
    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        "Рекомендации",
                        is_bold=True)
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}',
                        safe_get_value('recommendation'))

    for row in ws.iter_rows():
        for cell in row:
            if cell.value:
                if not cell.alignment or cell.alignment.wrap_text is not True:
                    cell.alignment = Alignment(
                        horizontal='left' if cell.column == 1 else 'left',
                        vertical='top',
                        wrap_text=True
                    )

    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            if cell.value:
                lines = str(cell.value).split('\n')
                for line in lines:
                    max_length = max(max_length, len(line))

        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    if save:
        wb.save(output_filename)
        return output_filename.split("testDjangosite")[1]
