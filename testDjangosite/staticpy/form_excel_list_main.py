from distutils.fancy_getopt import wrap_text

import openpyxl
from openpyxl.reader.excel import load_workbook
from openpyxl.styles import PatternFill, Alignment, Font, Border, Side
from openpyxl.workbook import Workbook
import pandas as pd

from testDjangosite.settings import BASE_DIR


def create_header(sheets):
    sheets['A1'].value = "Дата"
    sheets['C1'].value = "Номер"
    sheets['E1'].value = "Субъект РФ"
    sheets['G1'].value = "Лесничество"
    sheets['I1'].value = "Участковое лесничество"
    sheets['K1'].value = "Урочище"
    sheets['M1'].value = "Выдел"
    sheets.merge_cells("A1:B1")
    sheets.merge_cells("C1:D1")
    sheets.merge_cells("E1:F1")
    sheets.merge_cells("G1:H1")
    sheets.merge_cells("I1:J1")
    sheets.merge_cells("K1:L1")
    # sheets.merge_cells("A1:B1")


def form_getlistregion(data: dict ):
    filepath = f'{BASE_DIR}/media/excel_files/listregion/listregionfilters_{data["id_user"]}.xlsx'
    col = ["Номер", 'Дата', 'Участковое лесничество', 'лесничество', 'Субъект РФ', 'Урочище', 'Выдел', 'Площадь']
    col2 = ["Номер", 'Дата', 'Участковое лесничество', 'лесничество', 'Субъект РФ', 'Урочище', 'Выдел', 'Квартал']
    df = pd.DataFrame(data['data'])
    # field = loads(dumps(data['data_field']))
    df2 = pd.DataFrame(data['data_field'])
    df3 = pd.DataFrame(data['desc_field'])

    df.columns=col
    df2.columns=col2
    df3.columns=col2
    df.to_excel(filepath)
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name="Перечетная ведомость")
        df2.to_excel(writer, sheet_name="Полевая карточка")
        df3.to_excel(writer, sheet_name="Описание участка")

    wb = load_workbook(filepath)
    sheets = wb.active

    sheets_names = wb.sheetnames

    for j in range(0, 3):
        sheets = wb[sheets_names[j]]
        for i in range(7):
            rd = sheets.row_dimensions[i]
            rd.height = 30
        for row in sheets.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')

    wb.save(filepath)

    return filepath.split("testDjangosite")[1]

