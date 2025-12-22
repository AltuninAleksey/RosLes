import openpyxl
from openpyxl.styles import PatternFill, Alignment, Font, Border, Side
from openpyxl.workbook import Workbook
from openpyxl.utils import get_column_letter
from testDjangosite.settings import BASE_DIR


def create_forest_survey_excel(data, breeds_data = None, wb=None, save: bool = True):
    if save:
        output_filename = f'{BASE_DIR}/media/excel_files/field/field_{data["id"]}.xlsx'

    if wb is None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Полевая карточка"
    else:
        ws = wb.create_sheet("Полевая карточка")
    try:
        if 'не соответствует' in data.get('conclusion', ' ').lower():
            data['point7date'] = ''
            data['point7number'] = ''
            data['point7agreed'] = ''
            data['respond_farm'] = ''
    except:
        pass
    bold_font = Font(bold=True)
    border = Border(left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin'))
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    wrap_align = Alignment(horizontal='left', vertical='top', wrap_text=True)
    wrap_align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    wrap_align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    column_widths = {
        'A': 14,
        'B': 28,
        'C': 22,
        'D': 14,
        'E': 14,
        'F': 18,
        'G': 18,
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    def safe_get_value(value, default=''):
        if isinstance(value, dict):
            return str(value)
        elif value is None:
            return default
        return str(value)

    def create_wrapped_cell(ws, cell_range, value, is_bold=False, is_merged=True, alignment=wrap_align):
        if is_merged and ':' in cell_range:
            ws.merge_cells(cell_range)

        cell = ws[cell_range.split(':')[0]]
        cell.value = safe_get_value(value)

        if cell.alignment is None or not cell.alignment.wrap_text:
            cell.alignment = alignment

        if is_bold:
            cell.font = bold_font

        if value:
            text = safe_get_value(value)
            lines = text.split('\n')
            line_count = len(lines)

            for line in lines:
                if len(line) > 35:
                    line_count += len(line) // 35

            height = max(25, line_count * 18)
            row_num = cell.row
            ws.row_dimensions[row_num].height = height

        return cell

    def create_table_cell(ws, cell_address, value, is_header=False):
        cell = ws[cell_address]
        cell.value = safe_get_value(value)

        if is_header:
            cell.font = bold_font
            cell.border = border
            cell.alignment = wrap_align_center
        else:
            cell.border = border
            cell.alignment = wrap_align_center

        if value:
            text = safe_get_value(value)
            lines = text.split('\n')
            line_count = len(lines)

            for line in lines:
                if len(line) > 30:
                    line_count += len(line) // 30

            row_num = cell.row
            current_height = ws.row_dimensions[row_num].height

            if current_height is None:
                current_height = 15
            new_height = max(current_height, line_count * 18)
            ws.row_dimensions[row_num].height = max(25, new_height)

        return cell

    current_row = 1

    title_cell = create_wrapped_cell(ws, 'A1:G1', "ПОЛЕВАЯ КАРТОЧКА",
                                     is_bold=True, alignment=center_align)
    title_cell.font = Font(bold=True, size=14)

    subtitle_text = "натурного обследования с целью оценки качественных и количественных характеристик лесных насаждений при воспроизводстве лесов в рамках государственного мониторинга воспроизводства лесов"
    create_wrapped_cell(ws, 'A2:G2', subtitle_text, is_bold=True, alignment=center_align)

    current_row = 4

    basic_info = [
        ("участок №", data.get('number_region', '')),
        ("Субъект Российской Федерации", data.get('subject_rf', '')),
        ("Лесничество", data.get('forestly', '')),
        ("Лесной район", data.get('name_forest_district', '')),
        ("Участковое лесничество", data.get('district_forestly', '')),
        ("Урочище", data.get('name_dacha', '')),
        ("Квартал", data.get('name_quarter', '')),
        ("Выдел", data.get('soil_lot', '')),
        ("Площадь участка", data.get('sample_area', ''))
    ]

    for label, value in basic_info:
        ws[f'A{current_row}'] = label
        ws[f'A{current_row}'].font = bold_font
        ws[f'A{current_row}'].alignment = wrap_align_left
        ws[f'B{current_row}'] = safe_get_value(value)
        ws[f'B{current_row}'].alignment = wrap_align_left
        current_row += 1

    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}', "Характеристика участка",
                        is_bold=True, alignment=center_align)
    current_row += 1
    characteristics = [
        ("1.", "Целевое назначение лесов", data.get('purpose_of_forests', '')),
        ("", "Категория защитных лесов", data.get('forest_protection_category', '')),
        ("", "Особо защитные участки лесов", data.get('protected_areas_of_forests', '')),
        ("2.", "Участок находится в аренде (постоянном бессрочном пользовании)", "Да" if data.get('rent_area', '') == True else "Нет" ),
        ("3.", "Категория земель лесного фонда, на которой восстановлено лесное насаждение",
         data.get('category_of_forest_fund_lands', '')),
        ("4.", "Способ лесовосстановления", data.get('method_of_reforestation', '')),
        (
            "5.", "Срок проведения лесовосстановления (лесоразведения), месяц, год",
            data.get('time_of_reforestation', '')),
        ("6.", "Тип лесорастительных условий", data.get('forest_conditions', '')),
        ("", "Тип леса", data.get('forest_type', '')),
    ]

    for num, label, value in characteristics:
        ws[f'A{current_row}'] = num
        ws[f'A{current_row}'].alignment = wrap_align_left
        ws[f'B{current_row}'] = label
        ws[f'B{current_row}'].alignment = wrap_align_left
        if value:
            ws[f'C{current_row}'] = safe_get_value(value)
            ws[f'C{current_row}'].alignment = wrap_align_left
        current_row += 1

    point7_label = f"В {data.get('point7year', '')} году участок отнесен к землям, на которых расположены леса, по Акту отнесения земель, предназначенных для лесовосстановления к землям, на которых расположены леса от {data.get('point7date', '')} № {data.get('point7number', '')} утвержденному {data.get('point7agreed', '')}"
    ws[f'A{current_row}'] = "7."
    ws[f'A{current_row}'].alignment = wrap_align_left
    ws[f'B{current_row}'] = point7_label
    ws[f'B{current_row}'].alignment = wrap_align_left
    ws.merge_cells(f'B{current_row}:G{current_row}')
    current_row += 1

    point7_composition = f"{data.get('point7_natural_composition', '')}, {data.get('economy', '')} хозяйство, полнота (сомкнутость крон) {data.get('completeness', '')} запас {data.get('point7_stock', '')} м³/га"
    ws[f'A{current_row}'] = ""
    ws[f'A{current_row}'].alignment = wrap_align_left
    ws[f'B{current_row}'] = "Породный состав"
    ws[f'B{current_row}'].alignment = wrap_align_left
    ws[f'C{current_row}'] = point7_composition
    ws[f'C{current_row}'].alignment = wrap_align_left
    ws.merge_cells(f'C{current_row}:G{current_row}')
    current_row += 1

    current_row += 1

    headers = ["Коэффициент\nсостава", "Порода", "Возраст,\nлет", "Средняя\nвысота, м",
               "Средний\nдиаметр, см", "Количество учтенных\nдревесных растений,\nшт./га"]

    table1_start_row = current_row
    for col, header in enumerate(headers, 1):
        create_table_cell(ws, f'{get_column_letter(col)}{current_row}', header, is_header=True)

    current_row += 1

    coeff_data = data.get('coeff', [])
    for item in coeff_data:
        ratio_value = item.get('ratio', '')
        breed_value = item.get('breed', '')
        age_value = item.get('age', '')
        height_value = item.get('avg_height', '')
        diameter_value = item.get('avg_diametr', '')
        count_value = item.get('count_plants', '')

        create_table_cell(ws, f'A{current_row}', ratio_value)
        create_table_cell(ws, f'B{current_row}', breed_value)
        create_table_cell(ws, f'C{current_row}', age_value)
        create_table_cell(ws, f'D{current_row}', height_value)
        create_table_cell(ws, f'E{current_row}', diameter_value)
        create_table_cell(ws, f'F{current_row}', count_value)
        current_row += 1

    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}', "Результаты натурного обследования",
                        is_bold=True, alignment=center_align)
    current_row += 1

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}', "Закладка пробных площадей",
                        is_bold=True, alignment=center_align)
    current_row += 1

    survey_data = [
        ("8.", "Площадь 1 пробной площади", data.get('square_one_sample_area', ''), "кв. м"),
        ("9.", "Количество пробных площадей", data.get('count_sample_area', ''), "шт."),
    ]

    for num, label, value, unit in survey_data:
        ws[f'A{current_row}'] = num
        ws[f'A{current_row}'].alignment = wrap_align_left
        ws[f'B{current_row}'] = label
        ws[f'B{current_row}'].alignment = wrap_align_left
        if value:
            ws[f'C{current_row}'] = safe_get_value(value)
            ws[f'C{current_row}'].alignment = wrap_align_left
        if unit:
            ws[f'D{current_row}'] = unit
            ws[f'D{current_row}'].alignment = wrap_align_left
        current_row += 1

    # Пункт 10
    ws[f'A{current_row}'] = "10."
    ws[f'A{current_row}'].alignment = wrap_align_left
    ws[
        f'B{current_row}'] = "Координаты пробных площадей участка (в десятичных градусах с округлением до шестого знака):"
    ws[f'B{current_row}'].alignment = wrap_align_left
    ws.merge_cells(f'B{current_row}:G{current_row}')
    current_row += 1

    coord_headers = ["№ пробной\nплощади", "Широта", "Долгота"]

    table2_start_row = current_row
    for col, header in enumerate(coord_headers, 1):
        create_table_cell(ws, f'{get_column_letter(col)}{current_row}', header, is_header=True)

    current_row += 1

    samples_data = data.get('samples', [])
    for sample in samples_data:
        number_value = sample.get('number', '') if 'number' in sample else sample.get('number_sample', '')
        latitude_value = sample.get('latitude', '')
        longitude_value = sample.get('longitude', '')

        create_table_cell(ws, f'A{current_row}', number_value)
        create_table_cell(ws, f'B{current_row}', latitude_value)
        create_table_cell(ws, f'C{current_row}', longitude_value)
        current_row += 1

    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}', "Характеристика молодняка:",
                        is_bold=True, alignment=center_align)
    current_row += 1

    ws[f'A{current_row}'] = "11."
    ws[f'A{current_row}'].alignment = wrap_align_left
    ws[f'B{current_row}'] = "Породный состав"
    ws[f'B{current_row}'].alignment = wrap_align_left

    sapling_composition = f"{data.get('breed_composition', '')}, {data.get('economy_sapling', '')} хозяйство, полнота (сомкнутость крон) {data.get('completeness_sapling', '')}, запас {data.get('stock_sapling', '')} м³/га"
    ws[f'C{current_row}'] = sapling_composition
    ws[f'C{current_row}'].alignment = wrap_align_left
    ws.merge_cells(f'C{current_row}:G{current_row}')
    current_row += 2

    sapling_headers = ["Коэффициент\nсостава", "Порода", "Возраст,\nлет", "Средняя\nвысота, м",
                       "Средний\nдиаметр, см", "Количество учтенных\nдревесных растений,\nшт./га"]

    table3_start_row = current_row
    for col, header in enumerate(sapling_headers, 1):
        create_table_cell(ws, f'{get_column_letter(col)}{current_row}', header, is_header=True)

    current_row += 1

    saplings_data = data.get('saplings', [])
    for item in saplings_data:
        ratio_value = item.get('ratio_composition', '')
        if 'breed' in item:
            breed_value = item.get('breed', '')
        if 'id_breed' in item and not breeds_data is None:
            for breed in breeds_data['name_breeds']:
                if breed['id'] == item['id_breed']:
                    breed_value = breed['name_breed']
        age_value = item.get('age', '') if 'age' in item else item.get('avg_age', '')
        height_value = item.get('avg_height', '')
        diameter_value = item.get('diameter', '') if 'diameter' in item else item.get('avg_diameter', '')
        count_value = item.get('count_of_plants', '') if 'count_of_plants' in item else item.get('total', '')

        create_table_cell(ws, f'A{current_row}', ratio_value)
        create_table_cell(ws, f'B{current_row}', breed_value)
        create_table_cell(ws, f'C{current_row}', age_value)
        create_table_cell(ws, f'D{current_row}', height_value)
        create_table_cell(ws, f'E{current_row}', diameter_value)
        create_table_cell(ws, f'F{current_row}', count_value)
        current_row += 1

    current_row += 2

    create_wrapped_cell(ws, f'A{current_row}:G{current_row}', "Заключение",
                        is_bold=True, alignment=center_align)
    current_row += 1

    conclusion_data = [
        ("12.",
         f"Участок критериям и требованиям к молоднякам, площади которых подлежат отнесению к землям, на которых расположены леса, указанным в Правилах лесовосстановления, утвержденных приказом Минприроды России от {data.get('number_order', '')} (с изменениями) или лесохозяйственном регламенте лесничества:",
         data.get('conclusion', '')),
        ("",
         f"Участок хозяйству при отнесении к землям, на которых расположены леса, указанному в Акте отнесения земель, предназначенных для лесовосстановления, к землям, на которых расположены леса, от {data.get('point7date', '')} № {data.get('point7number', '')}, утвержденному {data.get('point7agreed', '')}",
         data.get('respond_farm', '')),
        ("13.",
         f"В случае несоответствия участка критериям и требованиям к молоднякам, площади которых подлежат отнесению к землям, на которых расположены леса, указанным в Правилах лесовосстановления, утвержденных приказом Минприроды России от {data.get('number_order', '')} (с изменениями) или лесохозяйственном регламенте лесничества, указать категорию земель лесного фонда, к которой относится участок",
         data.get('plot_farm_referring_land', '')),
        ("14.",
         "В случае отсутствия критериев для перевода в Правилах лесовосстановления необходимо внести реквизиты лесохозяйственного регламента:",
         data.get('details_regulations', '')),
        ("15.", "Рекомендации", data.get('recomendation', '')),
        ("16.", "Особенности участка", data.get('plot_features', ''))
    ]

    for num, label, value in conclusion_data:
        ws[f'A{current_row}'] = num
        ws[f'A{current_row}'].alignment = wrap_align_left
        ws[f'B{current_row}'] = label
        ws[f'B{current_row}'].alignment = wrap_align_left
        if value:
            ws[f'C{current_row}'] = safe_get_value(value)
            ws[f'C{current_row}'].alignment = wrap_align_left
            ws.merge_cells(f'C{current_row}:G{current_row}')
        current_row += 1

    current_row += 1

    survey_info = [
        ("Обследование провел", data.get('site_survey', '')),
        ("В присутствии", data.get('in_front', '')),
        ("Дата и время натурного обследования", data.get('date_and_time', '')),
        ("с", data.get('start_at', '')),
        ("по", data.get('end_at', ''))
    ]

    for label, value in survey_info:
        ws[f'A{current_row}'] = label
        ws[f'A{current_row}'].font = bold_font
        ws[f'A{current_row}'].alignment = wrap_align_left
        if value:
            ws[f'B{current_row}'] = safe_get_value(value)
            ws[f'B{current_row}'].alignment = wrap_align_left
        current_row += 1

    if save:
        wb.save(output_filename)
        return output_filename.split("testDjangosite")[1]

    return
