from distutils.fancy_getopt import wrap_text

import openpyxl
from openpyxl.styles import PatternFill, Alignment, Font, Border, Side
from openpyxl.workbook import Workbook
from djangoForest.models import *
from djangoForest.serializers import *
from testDjangosite.settings import BASE_DIR

# def merge_and_name(sheets, cur_col, name):
#     sheets.merge_cells(start_row=6, start_column=cur_col, end_row=6, end_column=cur_col)
#     cur_col = last_col
#     sheets.cell(row=6, column=last_col, value="Естественное возобновление (вегетативное)")

def set_border(ws, cell_range):
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))

    for row in ws[cell_range]:
        for cell in row:
            cell.border = thin_border


def create_header(sheets, data):

    # sheets = ws.active


    sheets['A2'].value = "П/П"
    sheets['A2'].border = Border(top=Side(style="thin"),bottom=Side(style="thin"))
    sheets['A2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['B2'].value = "Дата обследования"
    sheets['B2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['B2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    # sheets['C2'].value = "БПТ"
    # sheets['C2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    # sheets['C2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['C2'].value = "Лесничество"
    sheets['C2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['C2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['D2'].value = "Участковое лесничество"
    sheets['D2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['D2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['E2'].value = "Дача"
    sheets['E2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['E2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['F2'].value = "Квартал"
    sheets['F2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['F2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['G2'].value = "Выдел"
    sheets['G2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['G2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['H2'].value = "Площадь"
    sheets['H2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['H2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['I2'].value = "Акт №"
    sheets['I2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['I2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['J2'].value = "Дата по акту"
    sheets['J2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['J2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['K2'].value = "Год отнесения"
    sheets['K2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['K2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['L2'].value = "Аренда"
    sheets['L2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['L2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['M2'].value = "Категория"
    sheets['M2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['M2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['N2'].value = "Год проведения лесовосстановления"
    sheets['N2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['N2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['O2'].value = "Способ"
    sheets['O2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['O2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['P2'].value = "Метод"
    sheets['P2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['P2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['Q2'].value = "Соответствует критериям и требованиям"
    sheets['Q2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['Q2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    sheets['R2'].value = "Соответствует хозяйству"
    sheets['R2'].border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    sheets['R2'].border = Border(right=Side(style="thin"), bottom=Side(style="thin"))

    set_border(sheets, "A2:R2")



def main_draw(sheets, data: dict = None):

    create_header(sheets, data)

    methods = {2:"искусственное лесовосстановление", 3:"комбинированное лесовосстановление", 4: "естественное вследствие мер содействия лесовосстановлению", 5: "естественное вследсвие природных процессов"}
    cats = {2:"вырубка", 3:"гарь", 4: "прогалины и пустыри", 5: "погибшие насаждения", 6:'иное'}
    # data_test = ListRegionSerializer(ListRegion.objects.get(id=474)).data
    # print(data_test)
    print(data['data'])
    for row in range(len(data['data'])):
        first_sample = SampleSerializer(Sample.objects.filter(id_list_region=data['data'][row]['id']).first()).data
        field = FieldCardSerializer(FieldCard.objects.get(id_list_region=data['data'][row]['id'])).data
        desc = DescriptionRegionSerializer(DescriptionRegion.objects.get(id_list_region=data['data'][row]['id'])).data

        sheets.cell(row=row+3, column=1).value = data['data'][row]['number_region']
        sheets.cell(row=row+3, column=2).value = data['data'][row]['date']
        sheets.cell(row=row+3, column=3).value = data['data'][row]['forestly']
        sheets.cell(row=row + 3, column=4).value = data['data'][row]['id_district_forestly']
        sheets.cell(row=row + 3, column=5).value = data['data'][row]['dacha']
        sheets.cell(row=row + 3, column=6).value = data['data'][row]['name_quarter']
        sheets.cell(row=row + 3, column=7).value = data['data'][row]['soil_lot']
        sheets.cell(row=row + 3, column=8).value = data['data'][row]['sample_region']
        sheets.cell(row=row + 3, column=9).value = field['point7number']
        sheets.cell(row=row + 3, column=10).value = field['point7date']
        sheets.cell(row=row + 3, column=11).value = field['point7year']
        # sheets.cell(row=row + 3, column=13).value = desc['year_assignment_land']
        if field['rent_area'] is False:
            sheets.cell(row=row + 3, column=12).value = 'Нет'
        else:
            sheets.cell(row=row + 3, column=12).value = 'Да'
        if field['id_category_of_forest_fund_lands'] is None:
            sheets.cell(row=row + 3, column=13).value = "Не указано"
        else:
            sheets.cell(row=row + 3, column=13).value = cats[field['id_category_of_forest_fund_lands']]
        if field['id_method_of_reforestation'] is None:
            sheets.cell(row=row + 3, column=15).value = "Не указано"
        else:
            sheets.cell(row=row + 3, column=15).value = methods[field['id_method_of_reforestation']]
        sheets.cell(row=row + 3, column=14).value = field['time_of_reforestation']
        sheets.cell(row=row + 3, column=16).value = desc['recovery_method']
        if field['conclusion'] is None:
            sheets.cell(row=row + 3, column=17).value = 'Соответствует'
        else:
            sheets.cell(row=row + 3, column=17).value = field['conclusion']
        if field['respond_farm'] is False:
            sheets.cell(row=row + 3, column=18).value = "Нет"
        else:
            sheets.cell(row=row + 3, column=18).value ='Да'

        last_cell = sheets.cell(row=row+3, column=18)

    for col in sheets.columns:
        max_length = 0
        column = col[0].column_letter  # Get the column name
        for cell in col:
            try:  # Necessary to avoid error on empty cells
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        sheets.column_dimensions[column].width = adjusted_width

    set_border(sheets, f"A3:{last_cell.coordinate}")
    return last_cell